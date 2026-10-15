# -*- coding: utf-8 -*-
"""取文器：把各种格式的文件转成纯文本，供索引与“复制内容”使用。

全部离线实现。缺少可选依赖时降级为“只索引文件名”，不抛异常。
"""
from __future__ import annotations

import os

CODE_EXTS = {
    ".c", ".h", ".cc", ".cpp", ".hpp", ".cs", ".go", ".rs", ".rb", ".php",
    ".swift", ".m", ".mm", ".scala", ".java", ".kt", ".kts", ".ts", ".tsx",
    ".js", ".jsx", ".vue", ".svelte", ".sql", ".sh", ".bash", ".ps1", ".lua",
}


def sniff_decode(raw):
    """字节 -> 文本。识别 BOM、剔除二进制、多编码兜底。"""
    if not raw:
        return ""
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        try:
            return raw.decode("utf-16")
        except UnicodeDecodeError:
            pass
    if b"\x00" in raw[:8192]:
        return ""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def _read_bytes(path, guard=None):
    if guard is not None:
        with guard.open_source(path, "rb") as handle:
            return handle.read()
    with open(path, "rb") as handle:
        return handle.read()


def _docx_text(path):
    try:
        from docx import Document
    except ImportError:
        return "", "缺少 python-docx"
    try:
        document = Document(path)
    except Exception as exc:  # 损坏或加密文档
        return "", "docx 解析失败:%s" % type(exc).__name__
    parts = []
    for paragraph in document.paragraphs:
        if paragraph.text and paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts), ""


def _pdf_text(path, max_pages=40):
    try:
        from pypdf import PdfReader
    except ImportError:
        return "", "缺少 pypdf"
    try:
        reader = PdfReader(path)
    except Exception as exc:
        return "", "pdf 解析失败:%s" % type(exc).__name__
    chunks = []
    for index, page in enumerate(reader.pages):
        if index >= max_pages:
            break
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(chunks), ""


def extract_text(path, ext, max_chars=400000, guard=None, text_exts=None):
    """返回 (text, error, truncated)。

    error 为空串表示成功；truncated 为 True 表示正文被 max_chars 截断过 ——
    这个事实必须往上传，调用方有责任让用户看见（否则"后半部分搜不到"会被当成不存在）。
    """
    ext = (ext or os.path.splitext(path)[1]).lower()
    try:
        if ext == ".docx":
            text, error = _docx_text(path)
        elif ext == ".pdf":
            text, error = _pdf_text(path)
        elif ext in (".xlsx", ".xls"):
            return "", "表格仅索引文件名", False
        else:
            raw = _read_bytes(path, guard)
            text = sniff_decode(raw)
            error = "" if text else "二进制或无文本内容"
    except PermissionError:
        return "", "无读取权限", False
    except OSError as exc:
        return "", "读取失败:%s" % type(exc).__name__, False
    except Exception as exc:  # 兜底，任何取文异常都不应中断索引
        return "", "解析异常:%s" % type(exc).__name__, False

    truncated = bool(max_chars) and len(text) > max_chars
    if truncated:
        text = text[:max_chars]
    return text, error, truncated


def read_for_copy(path, max_chars=2000000, guard=None):
    """给"复制内容"用：返回 (text, error, truncated)。"""
    ext = os.path.splitext(path)[1].lower()
    return extract_text(path, ext, max_chars, guard=guard)


def read_for_search(path, guard=None):
    """给"文件内检索"用：直接只读源文件，返回 (kind, segments, error)。

    kind: 'text' | 'docx' | 'pdf' | 'jsonl' | 'binary'
    segments: [(序号, 文本), ...] —— 文本=行号，DOCX/PDF=第 N 段，JSONL=第 N 条消息。
    不能用 FTS 索引的 body：它已被 CJK 逐字切分、换行换成 ·、并截断，拿不到行号与原文。
    """
    ext = os.path.splitext(path)[1].lower()
    if ext in (".jsonl", ".ndjson"):
        raw = _read_bytes(path, guard)
        text = sniff_decode(raw)
        if not text:
            return "binary", [], "二进制或无文本内容"
        return "jsonl", list(enumerate(text.splitlines(), 1)), ""
    if ext == ".docx":
        text, error = _docx_text(path)
        if error:
            return "docx", [], error
        return "docx", list(enumerate(text.split("\n"), 1)), ""
    if ext == ".pdf":
        text, error = _pdf_text(path)
        if error:
            return "pdf", [], error
        paragraphs = [p.strip() for p in _split_paragraphs(text) if p.strip()]
        return "pdf", list(enumerate(paragraphs, 1)), ""
    raw = _read_bytes(path, guard)
    text = sniff_decode(raw)
    if not text:
        return "binary", [], "二进制或无文本内容"
    return "text", list(enumerate(text.splitlines(), 1)), ""


def _split_paragraphs(text):
    import re
    return re.split(r"\n\s*\n", text or "")
